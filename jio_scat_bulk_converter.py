#!/usr/bin/env python3
"""JioMart PDF invoice to SCAT XLSX bulk converter.

Run dashboard:
    streamlit run jio_scat_bulk_converter.py

Run command line:
    python jio_scat_bulk_converter.py --input ./pdfs --template "JIO Scat format.xlsx" --output SCAT_Output.xlsx
"""
from __future__ import annotations

import argparse
import io
import re
import subprocess
import tempfile
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import BinaryIO, Iterable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MONEY = Decimal("0.01")
SAC_ORDER = ("996211", "996819", "999794")
ITEM_NAMES = {
    "Seller Commission", "FixedFee", "CollectionFee", "ShippingCharge",
    "ShippingPenalty", "ReverseShippingCharge"
}

@dataclass
class Invoice:
    number: str
    invoice_date: date
    taxable_total: Decimal
    grand_total: Decimal
    sac_amounts: dict[str, Decimal] = field(default_factory=dict)
    source: str = ""
    warnings: list[str] = field(default_factory=list)


def money(value: str) -> Decimal:
    cleaned = re.sub(r"[^0-9.()-]", "", value or "").replace("(", "-").replace(")", "")
    return Decimal(cleaned or "0").quantize(MONEY, rounding=ROUND_HALF_UP)


def extract_text(pdf: bytes) -> str:
    """Extract text with pypdf, falling back to pdftotext when installed."""
    try:
        from pypdf import PdfReader
        text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(pdf)).pages)
        if len(text.strip()) > 200:
            return text
    except Exception:
        pass
    with tempfile.NamedTemporaryFile(suffix=".pdf") as src, tempfile.NamedTemporaryFile(suffix=".txt") as dst:
        src.write(pdf); src.flush()
        try:
            subprocess.run(["pdftotext", "-layout", src.name, dst.name], check=True, capture_output=True)
            return Path(dst.name).read_text(encoding="utf-8", errors="ignore")
        except Exception as exc:
            raise ValueError("Could not extract PDF text. Install pypdf or Poppler pdftotext.") from exc


def parse_invoice(pdf: bytes, source: str = "") -> Invoice:
    text = extract_text(pdf).replace("\xa0", " ")
    no_match = re.search(r"Invoice\s+Number\s*[-:]\s*([A-Z0-9/-]+)", text, re.I)
    date_match = re.search(r"Invoice\s+Date\s*[-:]\s*(\d{1,2}/\d{1,2}/\d{4})", text, re.I)
    total_match = re.search(r"(?m)^\s*Total\s+.*?₹?\s*([\d,]+\.\d{2})\s+.*?₹?\s*([\d,]+\.\d{2})\s+.*?₹?\s*([\d,]+\.\d{2})\s*$", text)
    grand_match = re.search(r"Grand\s+Total\s+₹?\s*([\d,]+\.\d{2})", text, re.I)
    if not no_match or not date_match:
        raise ValueError(f"{source or 'PDF'}: invoice number/date not found")

    invoice_date = datetime.strptime(date_match.group(1), "%d/%m/%Y").date()
    sac_amounts = {sac: Decimal("0.00") for sac in SAC_ORDER}

    # JioMart invoice item rows: item name, SAC (rate), taxable value, tax type...
    item_pattern = re.compile(
        r"(?m)^\s*(Seller Commission|FixedFee|CollectionFee|ShippingCharge|ShippingPenalty|ReverseShippingCharge)"
        r"\s+(\d{6})\s*\([^)]*\)\s+₹?\s*([\d,]+\.\d{2})"
    )
    for _item, sac, taxable in item_pattern.findall(text):
        if sac in sac_amounts:
            sac_amounts[sac] += money(taxable)

    if total_match:
        taxable_total = money(total_match.group(1))
        grand_total = money(total_match.group(3))
    else:
        taxable_total = sum(sac_amounts.values(), Decimal("0.00")).quantize(MONEY)
        if not grand_match:
            raise ValueError(f"{source or no_match.group(1)}: invoice totals not found")
        grand_total = money(grand_match.group(1))

    warnings = []
    parsed_sum = sum(sac_amounts.values(), Decimal("0.00")).quantize(MONEY)
    if parsed_sum != taxable_total:
        warnings.append(f"SAC total {parsed_sum} differs from taxable total {taxable_total}")
    if grand_total < taxable_total:
        warnings.append("Grand total is lower than taxable total")

    return Invoice(no_match.group(1), invoice_date, taxable_total, grand_total, sac_amounts, source, warnings)


def copy_row_style(ws, source_row: int, target_row: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, ws.max_column + 1):
        src, dst = ws.cell(source_row, col), ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.protection:
            dst.protection = copy(src.protection)


def month_text(d: date) -> tuple[str, str, str]:
    short = d.strftime("%b'%y")
    return f"SELLING EXPENSES_{short}", f"Selling Expenses for the month  {short}", d.replace(day=1).strftime("%d.%m.%Y")


def month_end(d: date) -> date:
    if d.month == 12:
        return date(d.year, 12, 31)
    return date(d.year, d.month + 1, 1) - __import__('datetime').timedelta(days=1)


def build_workbook(invoices: list[Invoice], template: str | Path | BinaryIO, posting_date: date | None = None) -> bytes:
    wb = load_workbook(template)
    ws = wb.active
    if ws.max_column < 20:
        raise ValueError("Template must contain the 20 SCAT columns.")

    # Capture the template's invoice-header and GL-line styles before removing sample rows.
    header_styles = [copy(ws.cell(2, c)._style) for c in range(1, 21)] if ws.max_row >= 2 else None
    line_styles = [copy(ws.cell(3, c)._style) for c in range(1, 21)] if ws.max_row >= 3 else header_styles
    header_height = ws.row_dimensions[2].height if ws.max_row >= 2 else None
    line_height = ws.row_dimensions[3].height if ws.max_row >= 3 else None

    # Keep header and remove sample data/merged ranges touching data rows.
    for merged in list(ws.merged_cells.ranges):
        if merged.min_row >= 2:
            ws.unmerge_cells(str(merged))
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    current = 2
    for sl_no, inv in enumerate(invoices, 1):
        if header_styles:
            for c, style in enumerate(header_styles, 1):
                ws.cell(current, c)._style = copy(style)
            ws.row_dimensions[current].height = header_height
        if line_styles:
            for r in range(current + 1, current + 4):
                for c, style in enumerate(line_styles, 1):
                    ws.cell(r, c)._style = copy(style)
                ws.row_dimensions[r].height = line_height
        header_text, line_text, _ = month_text(inv.invoice_date)
        value_date = month_end(inv.invoice_date)
        post = posting_date or date.today()

        values = [sl_no, 226222, inv.invoice_date, post, inv.number, float(inv.grand_total),
                  "MRA", "EMHO", header_text, "I6", "I6", float(inv.taxable_total)]
        for col, value in enumerate(values, 1):
            ws.cell(current, col, value)
        ws.cell(current, 3).number_format = "dd.mm.yyyy"
        ws.cell(current, 4).number_format = "dd.mm.yyyy"
        ws.cell(current, 6).number_format = "#,##0.00"
        ws.cell(current, 12).number_format = "#,##0.00"

        for offset, sac in enumerate(SAC_ORDER, 1):
            row = current + offset
            amount = inv.sac_amounts.get(sac, Decimal("0.00"))
            ws.cell(row, 1, sl_no)
            ws.cell(row, 13, 510484)
            ws.cell(row, 14, float(amount) if amount != 0 else None)
            ws.cell(row, 15, "XD")
            ws.cell(row, 16, "CORP91")
            ws.cell(row, 17, "OTHERS1000")
            ws.cell(row, 18, value_date)
            ws.cell(row, 19, line_text)
            ws.cell(row, 20, int(sac))
            ws.cell(row, 14).number_format = "#,##0.00"
            ws.cell(row, 18).number_format = "dd.mm.yyyy"
        current += 4

    ws.freeze_panes = "A2"
    widths = [8,12,14,14,22,14,16,14,26,10,10,14,14,14,10,14,16,14,38,16]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.auto_filter.ref = f"A1:T{max(1, ws.max_row)}"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    out = io.BytesIO(); wb.save(out); return out.getvalue()


def convert(pdf_files: Iterable[tuple[str, bytes]], template, posting_date: date | None = None):
    invoices, errors = [], []
    seen = set()
    for name, data in pdf_files:
        try:
            inv = parse_invoice(data, name)
            if inv.number in seen:
                errors.append(f"{name}: duplicate invoice {inv.number}")
                continue
            seen.add(inv.number); invoices.append(inv)
        except Exception as exc:
            errors.append(str(exc))
    invoices.sort(key=lambda x: (x.invoice_date, x.number))
    if not invoices:
        raise ValueError("No valid invoices were extracted. " + "; ".join(errors))
    return build_workbook(invoices, template, posting_date), invoices, errors


def streamlit_app() -> None:
    import streamlit as st
    st.set_page_config(page_title="JioMart PDF to SCAT", page_icon="📄", layout="wide")
    st.title("JioMart Invoice to SCAT Converter")
    st.caption("Upload multiple JioMart PDF invoices and the approved SCAT Excel template.")
    template = st.file_uploader("SCAT template (.xlsx)", type=["xlsx"])
    pdfs = st.file_uploader("Invoice PDFs", type=["pdf"], accept_multiple_files=True)
    posting_date = st.date_input("Posting date", value=date.today())
    if st.button("Generate SCAT Excel", type="primary", disabled=not (template and pdfs)):
        with st.spinner("Extracting and converting invoices..."):
            try:
                output, invoices, errors = convert(
                    [(f.name, f.getvalue()) for f in pdfs], io.BytesIO(template.getvalue()), posting_date
                )
                st.success(f"Processed {len(invoices)} invoice(s).")
                rows = [{"Invoice": i.number, "Date": i.invoice_date, "Taxable": float(i.taxable_total),
                         "Grand Total": float(i.grand_total), "Status": "Review" if i.warnings else "OK"}
                        for i in invoices]
                st.dataframe(rows, use_container_width=True, hide_index=True)
                for i in invoices:
                    for warning in i.warnings:
                        st.warning(f"{i.number}: {warning}")
                for error in errors:
                    st.error(error)
                st.download_button("Download SCAT XLSX", output, "JIO_SCAT_Bulk_Output.xlsx",
                                   "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            except Exception as exc:
                st.error(str(exc))


def cli() -> None:
    parser = argparse.ArgumentParser(description="Convert JioMart PDF invoices to SCAT XLSX")
    parser.add_argument("--input", help="Folder containing PDF invoices")
    parser.add_argument("--template", help="SCAT .xlsx template")
    parser.add_argument("--output", default="JIO_SCAT_Bulk_Output.xlsx")
    parser.add_argument("--posting-date", help="Posting date in DD.MM.YYYY")
    args = parser.parse_args()
    if not args.input or not args.template:
        streamlit_app(); return
    posting = datetime.strptime(args.posting_date, "%d.%m.%Y").date() if args.posting_date else date.today()
    files = [(p.name, p.read_bytes()) for p in sorted(Path(args.input).glob("*.pdf"))]
    output, invoices, errors = convert(files, args.template, posting)
    Path(args.output).write_bytes(output)
    print(f"Created {args.output} for {len(invoices)} invoice(s).")
    for error in errors: print("ERROR:", error)

if __name__ == "__main__":
    cli()
